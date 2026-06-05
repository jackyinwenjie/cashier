"""
从 test_data.json 生成 test_data.xlsx
运行: python gen_excel.py
"""
import json
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

HERE = Path(__file__).parent

# ── 读取 JSON ──
with open(HERE / "test_data.json", "r", encoding="utf-8") as f:
    TD = json.load(f)

VARS = TD.get("variables", {})

# ── 列定义 ──
HEADERS = ["id", "feature", "story", "title", "method", "path",
           "data", "assert", "extract", "special", "sign", "fixture"]

# ── 转换 $var → {{ var }} ──
def to_jinja(val):
    """递归将 $VAR → {{ VAR }}"""
    if isinstance(val, str):
        return re.sub(r'\$([A-Za-z_]\w*)', r'{{ \1 }}', val)
    if isinstance(val, list):
        return [to_jinja(v) for v in val]
    if isinstance(val, dict):
        return {k: to_jinja(v) for k, v in val.items()}
    return val


# ── 功能映射 ──
FEATURE_MAP = {
    "登录验证": "登录验证",
    "基础数据": "基础数据",
    "收银流程": "收银流程",
    "寄存支取": "寄存支取",
    "预订取消": "预订取消",
    "会员充值扣款": "会员充值扣款",
}

# ── 构建行 ──
rows = []
for suite in TD["suites"]:
    feature = FEATURE_MAP.get(suite["name"], suite["name"])
    tests = suite["tests"]
    for test in tests:
        data_obj = to_jinja(test.get("data", {}))
        assert_obj = test.get("assert", {})
        extract_obj = test.get("extract", {})
        special = test.get("special", "")
        sign_val = str(test.get("sign", "true")).lower()
        fixture = suite.get("fixture", "api")

        row = [
            str(test.get("id", "")),
            feature,
            suite["name"],
            test.get("desc", ""),
            test.get("method", "POST"),
            test.get("path", ""),
            json.dumps(data_obj, ensure_ascii=False),
            json.dumps(assert_obj, ensure_ascii=False),
            json.dumps(extract_obj, ensure_ascii=False),
            special,
            sign_val,
            fixture,
        ]
        rows.append(row)

# ── 写入 Excel ──
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "测试用例"

# 样式
header_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
data_align = Alignment(vertical="top", wrap_text=True)

# 写表头
for col_idx, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=col_idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_align

# 写数据
for row_idx, row_data in enumerate(rows, 2):
    for col_idx, val in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = data_align

# 设置列宽
col_widths = [25, 12, 12, 35, 8, 40, 55, 35, 35, 25, 8, 8]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

# 冻结首行
ws.freeze_panes = "A2"

# 自动筛选
ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(HEADERS))}1"

output_path = HERE / "test_data.xlsx"
wb.save(str(output_path))
print(f"Done: {output_path} ({len(rows)} cases)")
